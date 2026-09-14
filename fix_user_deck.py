"""
Targeted corrections to the hand-built SIH deck.

    python fix_user_deck.py

Edits ONLY claims that are not true of this repository, rewording each to what
is. Nothing is moved, resized or restyled; no layout is touched.

Two things checked and deliberately left alone:

  the footer bars were moved down by hand (7.06-7.22, not the template 6.95),
  so nothing actually overlapped them; and slide 5's footnote sits ON the bar
  on purpose - it is white text and reads perfectly there.

The original is backed up beside the file before anything is written.
"""
from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

DECK = Path(r"C:\Users\KARAN\Downloads\SIH2026-IDEA-Presentation-Format (1).pptx")
BACKUP = DECK.with_name(DECK.stem + " (original backup).pptx")
FOOTER_TOP = 6.95


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:                      # group
            yield from walk(sh.shapes)
        else:
            yield sh


def index(prs):
    out = {}
    for i, s in enumerate(prs.slides, 1):
        for sh in walk(s.shapes):
            out[(i, sh.name)] = sh
    return out


def set_text(shape, new):
    """Replace the text, keeping the first run's formatting and the box as is."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    if not p0.runs:
        tf.text = new
        return
    keep = p0.runs[0]
    keep.text = new
    for r in list(p0.runs[1:]):
        r._r.getparent().remove(r._r)
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)


def main():
    if not BACKUP.exists():
        shutil.copy2(DECK, BACKUP)
        print("backed up to %s" % BACKUP.name)

    prs = Presentation(str(BACKUP if BACKUP.exists() else DECK))
    ix = index(prs)
    done = []

    def edit(slide, name, new, why):
        sh = ix.get((slide, name))
        if sh is None or not sh.has_text_frame:
            print("  ! not found: slide %d %s" % (slide, name))
            return
        old = sh.text_frame.text.strip().replace("\n", " ")
        set_text(sh, new)
        done.append((slide, why, old, new))

    # ---------------------------------------------------------------- #
    # 1. claims that are not true of the code
    # ---------------------------------------------------------------- #
    # Slide 3 - the deployment strip is an architecture, not something built.
    # Only the first column exists today.
    edit(3, "TextBox 15", "DATA SOURCES · BUILT", "label the strip honestly")
    edit(3, "TextBox 21", "DATA LAYER · PLANNED", "label the strip honestly")
    edit(3, "TextBox 23", "API / SERVING · PLANNED", "label the strip honestly")
    edit(3, "TextBox 25", "CLIENTS · PLANNED", "label the strip honestly")

    # Slide 3 - stray edit in the team-name oval
    edit(3, "Oval 10", "Your Team Name", "the oval said SIH Winner")

    # Slide 4 - phone sensing was never built
    edit(4, "TextBox 27", "Validated against live running data",
         "phone sensing is not built")
    edit(4, "TextBox 28",
         "499 real arrivals collected from a live feed, compared with the "
         "simulator station by station.", "phone sensing is not built")

    # Slide 4 - shadow mode has not been done
    edit(4, "TextBox 48", "5  Shadow mode — planned", "shadow mode is not done")

    # Slide 4 - this risk is out of date: a live feed IS connected
    edit(4, "TextBox 53", "Delay model not yet fitted",
         "the old risk is out of date, a feed is connected")
    edit(4, "TextBox 54",
         "Validated on 499 live arrivals: the tail is too heavy. Refit as "
         "data accumulates.",
         "the old risk is out of date, a feed is connected")

    # Slide 4 - phone motion risk replaced by a real one
    edit(4, "TextBox 17415", "No public archive of past arrivals",
         "phone motion is not built")
    edit(4, "TextBox 17416",
         "None exists, so history is accumulated by polling a live feed — "
         "already collecting.", "phone motion is not built")

    # Slide 5 - neither of these is implemented
    edit(5, "TextBox 17561", "Position inferred between stations by block-graph replay",
         "phone sensing is not built")
    edit(5, "TextBox 17569", "Re-forecast every 30 s from the current state",
         "mid-run re-planning is not implemented")

    # ---------------------------------------------------------------- #
    # 2. content sitting under the blue footer bar
    # ---------------------------------------------------------------- #
    # NOTE: slide 5's footnote sits ON the blue bar on purpose - it is white
    # text and reads perfectly there. Moving it off the bar turned it white on
    # white. Left exactly where it was.

    target = DECK
    try:
        prs.save(str(DECK))
    except PermissionError:
        # the deck is open in PowerPoint, so write the corrected copy beside it
        target = DECK.with_name(DECK.stem + " CORRECTED.pptx")
        prs.save(str(target))
        print("original is open in PowerPoint - wrote %s" % target.name)

    print("\n%d edits:" % len(done))
    for sl, why, old, new in done:
        print("  slide %d  [%s]" % (sl, why))
        print("      was: %s" % old[:78])
        print("      now: %s" % new[:78])

    # verify against where each slide's footer bar ACTUALLY sits - several
    # were moved down by hand, so a fixed 6.95 would report false positives
    chk = Presentation(str(target))
    bad = []
    for i, s in enumerate(chk.slides, 1):
        bar = min([sh.top / 914400 for sh in s.shapes
                   if sh.shape_type == 1 and sh.width and sh.width > 12 * 914400
                   and sh.top and sh.top > 5 * 914400] or [FOOTER_TOP])
        for sh in walk(s.shapes):
            if sh.top is None or "Placeholder" in sh.name:
                continue
            if sh.name.startswith(("Rectangle 8", "Rectangle 9", "Rectangle 24")):
                continue
            if sh.has_text_frame and sh.text_frame.text.strip() == str(i):
                continue                       # the page number sits on the bar
            if (sh.top + (sh.height or 0)) / 914400 > bar + 0.02:
                bad.append((i, sh.name, round((sh.top + sh.height) / 914400, 2),
                            "bar at %.2f" % bar))
    print("\nstill under the footer bar: %s"
          % (bad if bad else "nothing"))


if __name__ == "__main__":
    main()
