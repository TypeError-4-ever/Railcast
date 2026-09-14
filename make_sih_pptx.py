"""
Build the submission deck in the official SIH 2026 IDEA template.

    python make_sih_pptx.py   ->  outputs/RAILCAST_SIH26028_IDEA.pptx

The template caps the deck at six slides including the title, so this is not the
nine-slide presentation deck: it is the submission. Content is fitted into the
template's own slides, keeping its branding, footers and section titles.

Every claim on these slides is either measured in this repository or labelled as
planned. What that ruled out, and why:

  networkx / igraph   not imported anywhere - the simulator is a custom heapq
                      event loop, so the deck says so
  MAPIE               not used - the conformal layer is our own CQR
  SHAP                not used - attribution is LightGBM split gain
  PyTorch Geometric   no GNN was built
  FastAPI, Redis,     none of the serving or interface stack exists yet; it is
  PostgreSQL, Docker, on the slide under "planned for deployment", never as a
  Kubernetes, React   capability we have
  phone motion, SMS,  not built
  IVR, shadow mode

The headline accuracy numbers are real measurements, but of a simulator whose
delay model is still provisional - the slides say that where the numbers appear.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from make_pptx import (AMBER, AMBER_BG, CHART, FAINT, INK, LINE, MONO, MUTE,
                       RED, RED_BG, SANS, SURF, TEAL, TEAL_BG, WHITE, arrow,
                       card, elbow, label, node, rect, text)

TEMPLATE = Path(r"C:\Users\KARAN\Downloads\SIH2026-IDEA-Presentation-Format.pptx")
OUT = Path("outputs/RAILCAST_SIH26028_IDEA.pptx")

TEAM_NAME = "[TEAM NAME]"
TEAM_ID = "[TEAM ID]"

# usable area inside the template chrome (logo top-right, footer bar at 6.95)
X0, X1 = 0.45, 12.88
Y0, Y1 = 1.32, 6.86
CW = X1 - X0


# --------------------------------------------------------------------------- #
def drop_slide(prs, index):
    lst = prs.slides._sldIdLst
    lst.remove(list(lst)[index])


def kill(shape):
    shape._element.getparent().remove(shape._element)


def find(slide, *needles, exact_name=None):
    for sh in slide.shapes:
        if exact_name and sh.name == exact_name:
            return sh
        if sh.has_text_frame:
            t = sh.text_frame.text.lower()
            if any(n.lower() in t for n in needles):
                return sh
    return None


def set_title(slide, s):
    for sh in slide.shapes:
        if sh.is_placeholder and sh.placeholder_format.idx == 0:
            tf = sh.text_frame
            tf.text = s
            for p in tf.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(26)
                    r.font.bold = True
            return sh
    return None


def set_team_oval(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and "your team name" in sh.text_frame.text.lower():
            tf = sh.text_frame
            tf.text = TEAM_NAME
            for p in tf.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for r in p.runs:
                    r.font.size = Pt(10)
                    r.font.bold = True
            return


def strip_prompt(slide, *needles):
    sh = find(slide, *needles)
    if sh is not None:
        kill(sh)


# --------------------------------------------------------------------------- #
def s1_title(slide):
    box = find(slide, "problem statement id")
    if box is None:
        return
    tf = box.text_frame
    rows = [("Problem Statement ID", "SIH26028"),
            ("Problem Statement Title", "Dynamic Forecast of Expected Time of "
                                        "Arrival (ETA) for Coaching Trains"),
            ("Theme", "Smart Automation"),
            ("PS Category", "Software"),
            ("Team ID", TEAM_ID),
            ("Team Name (Registered on portal)", TEAM_NAME)]
    tf.clear()
    for i, (k, v) in enumerate(rows):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(9)
        a = p.add_run()
        a.text = k + " – "
        a.font.size = Pt(13)
        a.font.bold = True
        a.font.name = SANS
        b = p.add_run()
        b.text = v
        b.font.size = Pt(13)
        b.font.name = SANS


def s2_solution(slide):
    set_title(slide, "RAILCAST — an arrival window, not a single number")
    strip_prompt(slide, "proposed solution (describe")
    set_team_oval(slide)

    # ---- problem ----
    label(slide, X0, Y0, 4, "The problem", color=RED)
    probs = ["Today's ETA repeats the delay a train already has and trusts "
             "timetable padding.",
             "It cannot see a conflict building ahead, so a held train surprises "
             "everyone behind it.",
             "One number, no confidence — nobody knows when it is safe to act."]
    for i, t in enumerate(probs):
        rect(slide, X0, Y0 + 0.30 + i * 0.72, 3.55, 0.64, fill=RED_BG, line=LINE, lw=0.8)
        rect(slide, X0, Y0 + 0.30 + i * 0.72, 0.03, 0.64, fill=RED)
        text(slide, X0 + 0.14, Y0 + 0.40 + i * 0.72, 3.28, 0.5, t, size=10,
             color=INK, spacing=1.15)

    label(slide, X0, Y0 + 2.58, 4, "How it works — one forecast", color=AMBER)
    steps = [("A", "Free-run time", "physics of each section"),
             ("B", "Conflict simulation", "advance every train on the block graph"),
             ("C", "Learned residual", "corrects what the rules cannot"),
             ("D", "Calibrated window", "audited 80% coverage")]
    for i, (k, t, b) in enumerate(steps):
        y = Y0 + 2.90 + i * 0.66
        rect(slide, X0, y, 3.55, 0.58, fill=WHITE, line=LINE, lw=0.8)
        rect(slide, X0, y, 0.30, 0.58, fill=AMBER_BG)
        text(slide, X0 + 0.09, y + 0.17, 0.25, 0.2, k, size=12, bold=True,
             color=AMBER, align=PP_ALIGN.CENTER)
        text(slide, X0 + 0.40, y + 0.09, 3.0, 0.2, t, size=10.5, bold=True)
        text(slide, X0 + 0.40, y + 0.30, 3.0, 0.2, b, size=9, color=MUTE)

    # ---- the decision the simulator makes ----
    cx, bw = 4.32, 3.80
    label(slide, cx, Y0, 4, "The decision the simulator resolves", color=AMBER)
    node(slide, "data", cx, Y0 + 0.30, bw, 0.46, "LIVE STATE",
         "position, delay, block occupancy")
    arrow(slide, cx + bw / 2, Y0 + 0.76, cx + bw / 2, Y0 + 0.96)
    node(slide, "decision", cx + 0.40, Y0 + 0.96, bw - 0.80, 0.90,
         "Block section ahead occupied?", size=10.5)
    text(slide, cx + bw / 2 + 0.04, Y0 + 1.88, 0.4, 0.18, "yes", size=8.5,
         bold=True, color=AMBER)
    arrow(slide, cx + bw / 2, Y0 + 1.86, cx + bw / 2, Y0 + 2.06)
    node(slide, "decision", cx + 0.40, Y0 + 2.06, bw - 0.80, 0.90,
         "Higher-priority train behind?", size=10.5)
    lx, rx2 = cx + 0.18, cx + bw - 0.18
    dy = Y0 + 2.51                                  # the diamond centre line
    node(slide, "process", cx, Y0 + 3.16, 1.78, 0.50, "Loop the slower train",
         None, size=9.5)
    node(slide, "process", cx + 2.02, Y0 + 3.16, 1.78, 0.50, "Queue behind block",
         None, size=9.5)
    elbow(slide, [(cx + 0.40, dy), (lx, dy), (lx, Y0 + 3.14)], head=True)
    elbow(slide, [(cx + bw - 0.40, dy), (rx2, dy), (rx2, Y0 + 3.14)], head=True)
    text(slide, lx + 0.06, dy - 0.20, 0.4, 0.18, "yes", size=8, bold=True, color=AMBER)
    text(slide, rx2 - 0.34, dy - 0.20, 0.4, 0.18, "no", size=8, bold=True, color=MUTE)
    # both branches rejoin before the window
    elbow(slide, [(cx + 0.89, Y0 + 3.66), (cx + 0.89, Y0 + 3.82),
                  (cx + bw / 2, Y0 + 3.82)], head=False)
    elbow(slide, [(cx + 2.91, Y0 + 3.66), (cx + 2.91, Y0 + 3.82),
                  (cx + bw / 2, Y0 + 3.82)], head=False)
    arrow(slide, cx + bw / 2, Y0 + 3.82, cx + bw / 2, Y0 + 3.98)
    node(slide, "terminator", cx + 0.20, Y0 + 3.98, bw - 0.40, 0.52,
         "ARRIVAL WINDOW", "+ the train that caused the hold", size=10.5)
    elbow(slide, [(cx + bw - 0.20, Y0 + 4.24), (cx + bw + 0.12, Y0 + 4.24),
                  (cx + bw + 0.12, Y0 + 0.53), (cx + bw, Y0 + 0.53)],
          color=AMBER, dashed=True)
    text(slide, cx + 0.20, Y0 + 4.60, bw, 0.2,
         "every actual arrival re-forecasts the rest of the run", size=8.5,
         color=AMBER, align=PP_ALIGN.CENTER)

    # ---- innovation ----
    ix = 8.46
    label(slide, ix, Y0, 4, "What makes it different", color=TEAL)
    inno = [("Forecasts the cause, not the clock",
             "Trains advanced over the real block-section graph until they "
             "conflict — so the ETA moves for a reason we can name."),
            ("States how sure it is, and audits it",
             "80% window, coverage checked per segment, horizon and time of day "
             "— not pooled once and asserted."),
            ("Names the blocking train",
             "Attribution separates a controller's loop from a physical queue "
             "behind an occupied block."),
            ("Checks itself against reality",
             "A live feed collects real arrivals; the model is re-fitted to them "
             "and reports where it still misses.")]
    for i, (t, b) in enumerate(inno):
        card(slide, ix, Y0 + 0.30 + i * 1.26, 4.42, 1.12, t, b, accent=TEAL,
             fill=TEAL_BG, tsize=11, bsize=9.2)
    text(slide, ix, Y0 + 5.36, 4.42, 0.2,
         "Working prototype — every stage above runs today.", size=9.5,
         bold=True, color=TEAL)


def s3_technical(slide):
    set_title(slide, "TECHNICAL APPROACH")
    strip_prompt(slide, "technologies to be used")
    set_team_oval(slide)

    label(slide, X0, Y0, 6, "Pipeline — built and running today")
    y, hh = Y0 + 0.28, 0.86
    node(slide, "data", X0, y, 2.05, hh, "1 · DATA IN",
         "real timetable, stations,\nobserved weather", size=10.5)
    arrow(slide, X0 + 2.06, y + hh / 2, X0 + 2.30, y + hh / 2)
    node(slide, "process", X0 + 2.36, y, 2.05, hh, "2 · BLOCK GRAPH",
         "202 stations, 1384 km,\nsectional speeds", size=10.5)
    arrow(slide, X0 + 4.42, y + hh / 2, X0 + 4.66, y + hh / 2)
    core = node(slide, "process", X0 + 4.72, y - 0.08, 2.60, hh + 0.16,
                "3 · FORECAST CORE", "A free-run · B conflict sim\nC residual · D conformal",
                size=10.5)
    core.line.color.rgb = AMBER
    core.fill.fore_color.rgb = AMBER_BG
    arrow(slide, X0 + 7.33, y + hh / 2, X0 + 7.57, y + hh / 2)
    node(slide, "decision", X0 + 7.63, y - 0.20, 2.10, hh + 0.40,
         "Coverage on target?", size=10.5)
    arrow(slide, X0 + 9.74, y + hh / 2, X0 + 9.98, y + hh / 2)
    text(slide, X0 + 9.76, y + hh / 2 - 0.22, 0.4, 0.18, "yes", size=8.5,
         bold=True, color=TEAL)
    node(slide, "terminator", X0 + 10.04, y, 2.39, hh, "4 · ARRIVAL WINDOW",
         "+ attribution, connection risk,\ncosted precedence", size=10.5)
    elbow(slide, [(X0 + 8.68, y + hh + 0.20), (X0 + 8.68, y + 1.36),
                  (X0 + 6.02, y + 1.36), (X0 + 6.02, y + hh + 0.08)],
          color=AMBER, dashed=True)
    text(slide, X0 + 8.74, y + 1.10, 0.4, 0.18, "no", size=8.5, bold=True, color=AMBER)
    text(slide, X0 + 6.20, y + 1.40, 3.0, 0.18,
         "re-fit the conformal quantiles for the group that missed", size=8.5,
         color=AMBER)

    # ---- stacks, honestly split ----
    sy = Y0 + 2.16
    label(slide, X0, sy, 6, "Built and measured — this repository", color=TEAL)
    built = [("Language and models", "Python 3 · LightGBM (free-run, quantile "
              "residual) · scikit-learn · NumPy / pandas"),
             ("Simulation", "Custom event-driven block-graph simulator "
              "(heapq priority queue) — no graph library"),
             ("Uncertainty", "Own conformalised quantile regression, Mondrian "
              "groups by horizon × segment × time of day"),
             ("Data adapters", "Indian Railways working timetable (CC0) · "
              "Open-Meteo ERA5 weather · RailRadar live feed"),
             ("Attribution", "LightGBM split gain, plus the blocking train and "
              "cause recorded by the simulator")]
    for i, (t, b) in enumerate(built):
        y2 = sy + 0.26 + i * 0.56
        rect(slide, X0, y2, 6.02, 0.52, fill=WHITE, line=LINE, lw=0.8)
        rect(slide, X0, y2, 0.03, 0.52, fill=TEAL)
        text(slide, X0 + 0.14, y2 + 0.07, 5.7, 0.2, t, size=10, bold=True)
        text(slide, X0 + 0.14, y2 + 0.28, 5.75, 0.24, b, size=8.8, color=MUTE,
             font=MONO)

    px = X0 + 6.40
    label(slide, px, sy, 6, "Planned for deployment — not built yet", color=FAINT)
    planned = [("Serving", "FastAPI · Redis Streams · PostgreSQL + PostGIS · "
                "Docker / Kubernetes, sharded by zone"),
               ("Passenger interface", "React · MapLibre GL for app and station "
                "displays"),
               ("Reaching every passenger", "SMS and IVR (139) in regional "
                "languages, for travellers without a smartphone"),
               ("Between-station sensing", "Opt-in passenger-phone motion to close "
                "the 40–60 minute reporting gap"),
               ("Operations", "Drift monitor, scheduled retraining, and a "
                "shadow-mode pilot beside the existing ETA")]
    for i, (t, b) in enumerate(planned):
        y2 = sy + 0.26 + i * 0.56
        rect(slide, px, y2, 6.02, 0.52, fill=SURF, line=LINE, lw=0.8)
        rect(slide, px, y2, 0.03, 0.52, fill=FAINT)
        text(slide, px + 0.14, y2 + 0.07, 5.7, 0.2, t, size=10, bold=True, color=MUTE)
        text(slide, px + 0.14, y2 + 0.28, 5.75, 0.24, b, size=8.8, color=FAINT,
             font=MONO)

    text(slide, X0, Y1 - 0.26, CW, 0.2,
         "The prototype runs end to end on one commodity laptop with no GPU: "
         "0.46 ms per station forecast, 0.32 s to re-forecast all 688 station "
         "arrivals on the corridor.", size=9, color=MUTE)


def s4_feasibility(slide):
    set_title(slide, "FEASIBILITY AND VIABILITY")
    strip_prompt(slide, "analysis of the feasibility")
    set_team_oval(slide)

    label(slide, X0, Y0, 4, "Why it is feasible", color=TEAL)
    feas = [("It already runs", "The whole pipeline executes end to end on a "
             "laptop — corridor, simulator, models, calibration, evaluation."),
            ("No new trackside hardware", "It consumes position data Indian "
             "Railways already produces, through a swappable adapter."),
            ("Cheap to serve", "Boosted trees, no GPU. 0.46 ms per station "
             "forecast on one commodity node."),
            ("Free, open inputs", "Timetable and station master are CC0; ERA5 "
             "weather needs no key."),
            ("Degrades safely", "With the learned layer off, the rule-based "
             "simulator still produces a forecast.")]
    for i, (t, b) in enumerate(feas):
        card(slide, X0, Y0 + 0.28 + i * 0.86, 3.92, 0.76, t, b, accent=TEAL,
             tsize=10.5, bsize=8.8)

    mx = X0 + 4.16
    label(slide, mx, Y0, 4, "Risks and how we retire them", color=RED)
    risks = [("The delay model is not yet fitted to reality",
              "Validated against 499 real arrivals: our tail is too heavy. The "
              "collector runs daily; refit as data accumulates."),
             ("No public archive of past arrivals",
              "None exists, so history is accumulated by polling a live feed — "
              "already collecting."),
             ("Error compounding over a 40-hour run",
              "Retired: each station is forecast in one shot from current state, "
              "never chained."),
             ("Coverage good on average, bad in a season",
              "Retired: calibration is conditioned on segment, horizon and time "
              "of day, and audited per cell."),
             ("Controllers distrust a black box",
              "Every forecast names the blocking train and section. Advisory "
              "only — the system issues no movement order.")]
    for i, (t, b) in enumerate(risks):
        card(slide, mx, Y0 + 0.28 + i * 0.86, 4.30, 0.76, t, b, accent=RED,
             fill=RED_BG, tsize=10.5, bsize=8.8)

    rx = mx + 4.54
    label(slide, rx, Y0, 4, "Status today", color=AMBER)
    done = [("Corridor from the real timetable", True),
            ("Block-graph conflict simulator", True),
            ("Four-stage forecast core", True),
            ("Calibrated 80% window, audited", True),
            ("Live feed connected, collecting", True),
            ("Validation against real arrivals", True),
            ("Delay model fitted to real data", False),
            ("Serving stack and passenger apps", False),
            ("Shadow-mode pilot on a corridor", False)]
    for i, (t, ok) in enumerate(done):
        y = Y0 + 0.28 + i * 0.46
        rect(slide, rx, y, 3.58, 0.40, fill=TEAL_BG if ok else SURF,
             line=LINE, lw=0.7)
        text(slide, rx + 0.12, y + 0.10, 0.3, 0.2, "✓" if ok else "○", size=11,
             bold=True, color=TEAL if ok else FAINT)
        text(slide, rx + 0.44, y + 0.11, 3.0, 0.2, t, size=9.5,
             color=INK if ok else MUTE)
    text(slide, rx, Y0 + 4.48, 3.58, 0.4,
         "Six of nine stages are built and measured. The remaining three are "
         "engineering, not research.", size=9, color=MUTE, spacing=1.15)

    label(slide, X0, Y1 - 0.92, 5, "Roadmap")
    steps = ["Collect live arrivals", "Refit the delay model", "Zonal pilot",
             "Serving stack", "Passenger channels", "Shadow mode, then rollout"]
    bw = (CW - 5 * 0.12) / 6
    for i, t in enumerate(steps):
        x = X0 + i * (bw + 0.12)
        last = i == 5
        rect(slide, x, Y1 - 0.62, bw, 0.46, fill=TEAL_BG if last else SURF,
             line=TEAL if last else LINE, lw=0.8)
        text(slide, x + 0.10, Y1 - 0.56, 0.4, 0.16, "%02d" % (i + 1), size=8,
             bold=True, font=MONO, color=AMBER)
        text(slide, x + 0.10, Y1 - 0.38, bw - 0.2, 0.2, t, size=8.5)
        if not last:
            arrow(slide, x + bw, Y1 - 0.39, x + bw + 0.10, Y1 - 0.39)


def s5_impact(slide):
    set_title(slide, "IMPACT AND BENEFITS")
    strip_prompt(slide, "potential impact on the target")
    set_team_oval(slide)

    mets = [("60%", "lower mean ETA error", "26.9 → 10.9 min", AMBER),
            ("62%", "lower P90 error", "77 → 29 min", AMBER),
            ("78.4%", "coverage achieved", "of the nominal 80% window", TEAL),
            ("40 h", "forecast horizon", "station by station", TEAL)]
    for i, (v, lab, sub, col) in enumerate(mets):
        x = X0 + i * 3.14
        rect(slide, x, Y0, 2.98, 1.00, fill=WHITE, line=LINE, lw=0.8)
        rect(slide, x, Y0, 2.98, 0.03, fill=col)
        text(slide, x + 0.14, Y0 + 0.13, 2.7, 0.34, v, size=23, bold=True, color=col)
        text(slide, x + 0.14, Y0 + 0.50, 2.7, 0.2, lab, size=10, bold=True)
        text(slide, x + 0.14, Y0 + 0.71, 2.7, 0.2, sub, size=8.8, color=MUTE)

    text(slide, X0, Y0 + 1.06, CW, 0.2,
         "Measured on 73,268 forecasts from held-out days, against carry-forward "
         "computed on exactly the same rows.", size=9, color=MUTE)

    slide.shapes.add_picture(str(CHART / "01_mae_by_horizon.png"), Inches(X0),
                             Inches(Y0 + 1.32), width=Inches(5.95))
    text(slide, X0, Y0 + 4.36, 5.95, 0.44,
         "Carry-forward is adequate for the next station and collapses beyond a "
         "few hours — the horizon passengers actually plan against.", size=9,
         color=MUTE, spacing=1.15)

    bx = X0 + 6.28
    label(slide, bx, Y0 + 1.32, 6, "Who feels the difference", color=TEAL)
    who = [("Passengers", "An arrival window to plan around, the odds of making a "
            "connection, and when to leave home."),
           ("Section controllers", "Conflicts surfaced before they happen, with the "
            "blocking train named and the cost of each precedence call."),
           ("Station operations", "Platforms, cleaning turnaround and crew duty "
            "planned against a forecast with a stated confidence."),
           ("Feeder and logistics", "Bus, taxi and parcel operators schedule to the "
            "real arrival instead of idling at the kerb.")]
    for i, (t, b) in enumerate(who):
        card(slide, bx, Y0 + 1.62 + i * 0.80, 6.15, 0.70, t, b, accent=TEAL,
             fill=TEAL_BG, tsize=10.5, bsize=8.8)

    rect(slide, bx, Y0 + 4.90, 6.15, 0.62, fill=AMBER_BG, line=AMBER, lw=0.9)
    text(slide, bx + 0.14, Y0 + 4.98, 5.9, 0.5,
         [[("Honest scope. ", {"bold": True, "color": INK}),
           ("These are prototype results on a simulated corridor built from the "
            "real timetable. Delay generation is modelled, and validation against "
            "499 real arrivals shows it is not yet right — which is why the "
            "calibration loop exists.", {})]], size=8.8, color=MUTE, spacing=1.15)


def s6_references(slide):
    set_title(slide, "RESEARCH AND REFERENCES")
    strip_prompt(slide, "details / links of the reference")
    set_team_oval(slide)

    label(slide, X0, Y0, 5, "Data actually used", color=TEAL)
    srcs = [("Indian Railways working timetable", "datameet/railways, CC0 — 8,697 "
             "stations, 5,208 trains, 417,080 scheduled calls"),
            ("Weather", "ERA5 reanalysis via Open-Meteo — observed daily, per "
             "corridor segment, no API key"),
            ("Live running position", "RailRadar REST API — 499 observed arrivals "
             "collected across 10 trains so far"),
            ("Route geometry", "Station coordinates scaled to the published route "
             "distance; 1384 km reproduced exactly")]
    for i, (t, b) in enumerate(srcs):
        y = Y0 + 0.30 + i * 0.80
        rect(slide, X0, y, 6.05, 0.70, fill=WHITE, line=LINE, lw=0.8)
        rect(slide, X0, y, 0.03, 0.70, fill=TEAL)
        text(slide, X0 + 0.14, y + 0.09, 5.7, 0.2, t, size=10.5, bold=True)
        text(slide, X0 + 0.14, y + 0.31, 5.75, 0.34, b, size=9, color=MUTE,
             spacing=1.12)

    label(slide, X0, Y0 + 3.60, 5, "Scale of deployment", color=AMBER)
    scale = [("Prototype", "New Delhi – Mumbai Central, 1384 km, 202 stations, "
              "44 trains — built and evaluated"),
             ("Zonal pilot", "One zone, one control office, one model shard"),
             ("National", "Order of 13,000 coaching trains a day (Indian Railways "
              "published figure — confirm before submission)")]
    for i, (t, b) in enumerate(scale):
        y = Y0 + 3.90 + i * 0.56
        text(slide, X0, y, 1.55, 0.2, t, size=10, bold=True, color=AMBER)
        text(slide, X0 + 1.62, y, 4.45, 0.44, b, size=9, color=MUTE, spacing=1.12)

    rx = X0 + 6.40
    label(slide, rx, Y0, 5, "Research we build on")
    refs = ["R. Oneto et al., Train delay prediction systems: a big data "
            "analytics perspective, Big Data Research 11, 2018.",
            "P. Kecman, R. M. P. Goverde, Predictive modelling of running and "
            "dwell times in railway traffic, Public Transport 7, 2015.",
            "F. Corman, P. Kecman, Stochastic prediction of train delays in real "
            "time using Bayesian networks, Transp. Research Part C 95, 2018.",
            "Y. Romano, E. Patterson, E. Candès, Conformalized quantile "
            "regression, NeurIPS 32, 2019.",
            "V. Vovk, A. Gammerman, G. Shafer, Algorithmic Learning in a Random "
            "World, Springer 2005 — Mondrian conditional validity.",
            "G. Ke et al., LightGBM: a highly efficient gradient boosting "
            "decision tree, NeurIPS 30, 2017.",
            "The Digital Personal Data Protection Act, 2023 (India)."]
    for i, r in enumerate(refs):
        y = Y0 + 0.30 + i * 0.44
        text(slide, rx, y, 0.24, 0.2, "%d." % (i + 1), size=9, bold=True, color=AMBER)
        text(slide, rx + 0.28, y, 5.72, 0.4, r, size=9, color=MUTE, spacing=1.12)

    rect(slide, rx, Y0 + 3.44, 6.00, 0.92, fill=AMBER_BG, line=AMBER, lw=0.9)
    text(slide, rx + 0.14, Y0 + 3.54, 5.72, 0.2, "WHERE WE ADD TO THIS", size=8.5,
         bold=True, font=MONO, color=AMBER)
    text(slide, rx + 0.14, Y0 + 3.76, 5.72, 0.54,
         "None of the above combines a conflict-resolution simulator over the "
         "real block-section graph, a direct multi-horizon learned residual and "
         "conditionally calibrated intervals — then validates the result against "
         "live Indian Railways running data.", size=9, spacing=1.15)

    text(slide, rx, Y0 + 4.52, 6.00, 0.5,
         [[("Project links  ", {"bold": True, "color": INK}),
           ("code and reproducible results · simulation videos · evaluation "
            "tables and charts  —  [REPOSITORY LINK]", {})]], size=9, color=MUTE,
         spacing=1.15)


# --------------------------------------------------------------------------- #
def main():
    prs = Presentation(str(TEMPLATE))
    drop_slide(prs, 6)                       # the instructions slide
    slides = list(prs.slides)
    s1_title(slides[0])
    s2_solution(slides[1])
    s3_technical(slides[2])
    s4_feasibility(slides[3])
    s5_impact(slides[4])
    s6_references(slides[5])
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print("wrote %s  (%d slides, %.1f MB)"
          % (OUT, len(prs.slides._sldIdLst), OUT.stat().st_size / 1048576))


if __name__ == "__main__":
    main()
